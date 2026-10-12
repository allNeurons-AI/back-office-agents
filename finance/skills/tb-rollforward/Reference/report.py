"""report.py - write the trial-balance workbook (Summary, Consolidated, roll-ups, one tab per entity, Checks)."""
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from engine import totals

NUM = '#,##0.00;(#,##0.00);"-"'
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
F = "Arial"
thin = Side(style="thin", color="999999")


def lab(ctx, e):
    """Entity label: display name if the entities table has one, else the id."""
    return (ctx["ents"].get(e, {}) or {}).get("name") or str(e)


def _font(**k):
    return Font(name=F, size=k.pop("size", 10), **k)


def _sheet_name(wb, base, used):
    n = re.sub(r"[\[\]:*?/\\]", "-", str(base))[:31] or "Sheet"
    i, cand = 2, n
    while cand.lower() in used:
        cand = (n[:28] + "_%d" % i)
        i += 1
    used.add(cand.lower())
    return cand


def _header(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = _font(bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = _font(bold=True, size=14)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = _font(italic=True, color="555555")


def write_workbook(path, meta, tr, ctx, roll, checks, warnings, drift):
    wb = Workbook()
    used = set()
    rep = ctx["rep"]
    ents = sorted(tr, key=lambda e: lab(ctx, e).lower())
    basis = meta["basis"]

    # ---------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = _sheet_name(wb, "Summary", used)
    _title(ws, "%s - Trial balance by entity" % (meta.get("company") or "Company"),
           "As of %s  |  %s  |  reporting currency %s  |  run %s (%s)" % (
               meta["asof"], basis, rep, meta["run_time"], meta["run_type"]))
    r = 4
    if basis == "PRELIMINARY":
        ws.cell(row=r, column=1, value="PRELIMINARY: includes months that are not closed yet. Numbers can still change.").font = _font(bold=True, color="9C5700")
        ws.cell(row=r, column=1).fill = WARN_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        r += 2
    _header(ws, r, ["Entity", "Currency", "Total assets", "Total liabilities", "Total equity (incl. translation adj.)",
                    "Net income (profit +)", "Balance check"], [34, 10, 18, 18, 24, 20, 14])
    ws.row_dimensions[r].height = 32
    r += 1
    first = r
    gt = {"assets": 0, "liabilities": 0, "equity": 0, "net_income": 0}
    for e in ents:
        t = totals(tr[e])
        vals = [lab(ctx, e), ctx["ent_ccy"].get(e, rep), t["assets"], t["liabilities"], t["equity"], t["net_income"], t["check"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = _font()
            if i >= 3:
                c.number_format = NUM
        for k in gt:
            gt[k] += t[k]
        r += 1
    vals = ["Consolidated (%s)" % rep, rep, gt["assets"], gt["liabilities"], gt["equity"], gt["net_income"],
            sum(totals(tr[e])["check"] for e in ents)]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = _font(bold=True)
        c.fill = SUB_FILL
        c.border = Border(top=thin)
        if i >= 3:
            c.number_format = NUM
    r += 2
    ws.cell(row=r, column=1, value="Liabilities, equity and net income are shown as positive numbers for readability "
                                    "(debits are positive inside the tabs, credits negative).").font = _font(italic=True, color="555555")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    r += 2
    ws.cell(row=r, column=1, value="Run details").font = _font(bold=True, size=11)
    r += 1
    for k, v in meta["details"]:
        ws.cell(row=r, column=1, value=k).font = _font(bold=True)
        ws.cell(row=r, column=2, value=str(v)).font = _font()
        r += 1
    if warnings or drift:
        r += 1
        ws.cell(row=r, column=1, value="Warnings").font = _font(bold=True, size=11, color="9C0006")
        r += 1
        for w in warnings + ["DRIFT: %s %s %s - %s" % (d["kind"], d["period"], d["entity"], d["detail"]) for d in drift]:
            c = ws.cell(row=r, column=1, value=w)
            c.font = _font()
            c.fill = WARN_FILL
            r += 1

    # ---------------------------------------------------------------- Consolidated + roll-ups
    def matrix(title, names, sheet_label):
        w = wb.create_sheet(_sheet_name(wb, sheet_label, used))
        _title(w, title, "Reporting currency %s, as of %s (%s)" % (rep, meta["asof"], basis))
        accts = {}
        for e in names:
            for row in tr[e]:
                accts.setdefault(row["account"], row)
        order = sorted(accts, key=lambda a: (a == "CTA", len(a), a))
        _header(w, 4, ["Account", "Name", "Type"] + [lab(ctx, n) for n in names] + ["Total"], [12, 36, 16] + [18] * (len(names) + 1))
        w.freeze_panes = "D5"
        rr = 5
        for a in order:
            w.cell(row=rr, column=1, value=a).font = _font()
            w.cell(row=rr, column=2, value=accts[a]["name"]).font = _font()
            w.cell(row=rr, column=3, value=accts[a]["cls"].replace("_", " ")).font = _font()
            tot = 0.0
            for j, e in enumerate(names):
                v = next((x["tr"] for x in tr[e] if x["account"] == a), 0.0)
                tot += v
                c = w.cell(row=rr, column=4 + j, value=v)
                c.number_format = NUM
                c.font = _font()
            c = w.cell(row=rr, column=4 + len(names), value=tot)
            c.number_format = NUM
            c.font = _font(bold=True)
            rr += 1
        rr += 1
        for label, key in (("Total assets", "assets"), ("Total liabilities", "liabilities"),
                           ("Total equity (incl. translation adj.)", "equity"),
                           ("Net income (profit +)", "net_income"), ("Check: all rows add to zero", "check")):
            w.cell(row=rr, column=2, value=label).font = _font(bold=True)
            tot = 0.0
            for j, e in enumerate(names):
                v = totals(tr[e])[key]
                tot += v
                c = w.cell(row=rr, column=4 + j, value=v)
                c.number_format = NUM
                c.font = _font(bold=True)
                c.fill = SUB_FILL
            c = w.cell(row=rr, column=4 + len(names), value=tot)
            c.number_format = NUM
            c.font = _font(bold=True)
            c.fill = SUB_FILL
            rr += 1

    matrix("Consolidated", ents, "Consolidated")
    for node, leaves in roll.items():
        if set(leaves) == set(ents):
            continue
        matrix("Roll-up: %s" % lab(ctx, node), leaves, "Rollup - %s" % lab(ctx, node))

    # ---------------------------------------------------------------- Entity tabs
    for e in ents:
        w = wb.create_sheet(_sheet_name(wb, lab(ctx, e), used))
        ccy = ctx["ent_ccy"].get(e, rep)
        _title(w, "%s (%s)" % (lab(ctx, e), ccy), "As of %s (%s). Debit = positive, credit = negative." % (meta["asof"], basis))
        w["A3"] = "Parent: %s" % (lab(ctx, ctx["parent"].get(e)) if ctx["parent"].get(e) else "-")
        w["A3"].font = _font(color="555555")
        _header(w, 5, ["Account", "Name", "Type", "Balance (%s)" % ccy, "Translated how", "Closing rate",
                       "Debit (%s)" % rep, "Credit (%s)" % rep, "Net (%s)" % rep],
                [12, 36, 16, 18, 32, 12, 18, 18, 18])
        w.freeze_panes = "A6"
        rr = 6
        for row in tr[e]:
            vals = [row["account"], row["name"], row["cls"].replace("_", " "), row["local"], row["basis"],
                    row["rate"], max(row["tr"], 0), max(-row["tr"], 0), row["tr"]]
            for i, v in enumerate(vals, 1):
                c = w.cell(row=rr, column=i, value=v)
                c.font = _font()
                if i in (4, 7, 8, 9):
                    c.number_format = NUM
                if i == 6 and v is not None:
                    c.number_format = "0.000000"
            rr += 1
        t = totals(tr[e])
        rr += 1
        tot_d = sum(max(x["tr"], 0) for x in tr[e])
        tot_c = sum(max(-x["tr"], 0) for x in tr[e])
        for i, v in ((2, "Total"), (7, tot_d), (8, tot_c), (9, tot_d - tot_c)):
            c = w.cell(row=rr, column=i, value=v)
            c.font = _font(bold=True)
            c.fill = SUB_FILL
            if i >= 7:
                c.number_format = NUM
        rr += 2
        for label, key in (("Total assets", "assets"), ("Total liabilities", "liabilities"),
                           ("Total equity (incl. translation adj.)", "equity"), ("Net income (profit +)", "net_income")):
            w.cell(row=rr, column=2, value=label).font = _font(bold=True)
            c = w.cell(row=rr, column=9, value=t[key])
            c.number_format = NUM
            c.font = _font(bold=True)
            rr += 1

    # ---------------------------------------------------------------- Checks
    w = wb.create_sheet(_sheet_name(wb, "Checks", used))
    _title(w, "Checks")
    _header(w, 3, ["Check", "Scope", "Result", "Detail"], [48, 30, 10, 70])
    rr = 4
    for ck in checks:
        vals = [ck["name"], ck["scope"], ck["status"].upper(), ck["detail"]]
        for i, v in enumerate(vals, 1):
            c = w.cell(row=rr, column=i, value=v)
            c.font = _font()
            if i == 3:
                c.fill = {"ok": OK_FILL, "warn": WARN_FILL, "fail": BAD_FILL}[ck["status"]]
        rr += 1
    wb.save(path)
    return path
